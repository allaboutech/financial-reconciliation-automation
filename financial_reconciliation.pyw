import csv
import hashlib
import json
import os
import tkinter as tk
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from tkinter import filedialog, messagebox, ttk

try:
    from openpyxl import load_workbook, Workbook
except ImportError:
    load_workbook = Workbook = None

COLUMNS = ('transaction_id', 'account', 'date', 'amount')

def money(value):
    try:
        return Decimal(str(value).replace(',', '').replace('$', '').strip()).quantize(Decimal('0.01'))
    except (InvalidOperation, ValueError):
        raise ValueError(f'Invalid amount: {value!r}')

def load_records(path):
    if path.lower().endswith('.csv'):
        with open(path, newline='', encoding='utf-8-sig') as f:
            records = list(csv.DictReader(f))
    elif path.lower().endswith('.xlsx'):
        if load_workbook is None:
            raise RuntimeError('Install openpyxl to read Excel files: pip install openpyxl')
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = wb.active.iter_rows(values_only=True)
            headers = [str(x).strip().lower() if x is not None else '' for x in next(rows)]
            records = [dict(zip(headers, row)) for row in rows if any(x is not None for x in row)]
        finally:
            wb.close()
    else:
        raise ValueError('Choose a .csv or .xlsx file')
    result = []
    for line, record in enumerate(records, 2):
        record = {str(k).strip().lower(): v for k, v in record.items() if k is not None}
        missing = [c for c in COLUMNS if c not in record]
        if missing:
            raise ValueError('Missing columns: ' + ', '.join(missing))
        tid = str(record['transaction_id'] or '').strip()
        acct = str(record['account'] or '').strip()
        if not tid or not acct:
            raise ValueError(f'Row {line}: transaction_id and account are required')
        result.append({'transaction_id': tid, 'account': acct, 'date': str(record['date'] or '').strip(), 'amount': money(record['amount']), 'source_row': line})
    return result

def reconcile(gl, sub):
    left, right = defaultdict(list), defaultdict(list)
    for row in gl:
        left[(row['transaction_id'], row['account'])].append(row)
    for row in sub:
        right[(row['transaction_id'], row['account'])].append(row)
    exceptions = []
    matched = 0
    for key in sorted(set(left) | set(right)):
        a, b = left[key], right[key]
        if len(a) > 1 or len(b) > 1:
            exceptions.append((key[0], key[1], 'DUPLICATE_KEY', len(a), len(b), str(sum((r['amount'] for r in a), Decimal(0))), str(sum((r['amount'] for r in b), Decimal(0)))))
        elif not a:
            exceptions.append((key[0], key[1], 'MISSING_GL', 0, 1, '0.00', str(b[0]['amount'])))
        elif not b:
            exceptions.append((key[0], key[1], 'MISSING_SUBLEDGER', 1, 0, str(a[0]['amount']), '0.00'))
        elif a[0]['amount'] != b[0]['amount']:
            exceptions.append((key[0], key[1], 'AMOUNT_MISMATCH', 1, 1, str(a[0]['amount']), str(b[0]['amount'])))
        elif a[0]['date'] != b[0]['date']:
            exceptions.append((key[0], key[1], 'DATE_MISMATCH', 1, 1, str(a[0]['amount']), str(b[0]['amount'])))
        else:
            matched += 1
    return matched, exceptions

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title('Financial Reconciliation Automation | Controller Toolkit')
        self.geometry('1080x690')
        self.minsize(850, 550)
        self.paths = {'GL': tk.StringVar(), 'Subledger': tk.StringVar()}
        self.results = None
        ttk.Label(self, text='Financial Reconciliation Automation', font=('Segoe UI', 19, 'bold')).pack(anchor='w', padx=20, pady=(18, 3))
        ttk.Label(self, text='GL vs Subledger | Exception Detection | Audit-Ready Exports').pack(anchor='w', padx=20, pady=(0, 16))
        for name in ('GL', 'Subledger'):
            line = ttk.Frame(self)
            line.pack(fill='x', padx=20, pady=5)
            ttk.Label(line, text=name, width=12).pack(side='left')
            ttk.Entry(line, textvariable=self.paths[name]).pack(side='left', fill='x', expand=True, padx=8)
            ttk.Button(line, text='Browse Excel / CSV', command=lambda n=name: self.browse(n)).pack(side='left')
        controls = ttk.Frame(self)
        controls.pack(fill='x', padx=20, pady=15)
        ttk.Button(controls, text='Run Reconciliation', command=self.run).pack(side='left')
        ttk.Button(controls, text='Export Report', command=self.export).pack(side='left', padx=10)
        self.status = tk.StringVar(value='Ready. Required columns: transaction_id, account, date, amount')
        ttk.Label(self, textvariable=self.status, wraplength=1000).pack(anchor='w', padx=20, pady=5)
        headers = ('transaction_id', 'account', 'issue', 'gl_count', 'sub_count', 'gl_amount', 'sub_amount')
        area = ttk.Frame(self)
        area.pack(fill='both', expand=True, padx=20, pady=(5, 20))
        self.tree = ttk.Treeview(area, columns=headers, show='headings')
        for h in headers:
            self.tree.heading(h, text=h.replace('_', ' ').title())
            self.tree.column(h, width=130, anchor='w')
        yscroll = ttk.Scrollbar(area, orient='vertical', command=self.tree.yview)
        xscroll = ttk.Scrollbar(area, orient='horizontal', command=self.tree.xview)
        self.tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        self.tree.grid(row=0, column=0, sticky='nsew')
        yscroll.grid(row=0, column=1, sticky='ns')
        xscroll.grid(row=1, column=0, sticky='ew')
        area.rowconfigure(0, weight=1)
        area.columnconfigure(0, weight=1)
    def browse(self, name):
        path = filedialog.askopenfilename(filetypes=[('Excel and CSV', '*.xlsx *.csv')])
        if path:
            self.paths[name].set(path)
    def run(self):
        try:
            gl = load_records(self.paths['GL'].get())
            sub = load_records(self.paths['Subledger'].get())
            matched, exceptions = reconcile(gl, sub)
            self.results = (matched, exceptions, len(gl), len(sub))
            for item in self.tree.get_children():
                self.tree.delete(item)
            for row in exceptions:
                self.tree.insert('', 'end', values=row)
            self.status.set(f'GL rows: {len(gl):,} | Subledger rows: {len(sub):,} | Matched unique keys: {matched:,} | Exceptions: {len(exceptions):,}')
        except Exception as exc:
            messagebox.showerror('Reconciliation error', str(exc))
    def export(self):
        if self.results is None:
            messagebox.showwarning('No results', 'Run reconciliation first.')
            return
        path = filedialog.asksaveasfilename(defaultextension='.xlsx' if Workbook else '.csv', filetypes=[('Excel', '*.xlsx')] if Workbook else [('CSV', '*.csv')])
        if not path:
            return
        matched, exceptions, gl_count, sub_count = self.results
        headers = ['transaction_id', 'account', 'issue', 'gl_count', 'sub_count', 'gl_amount', 'sub_amount']
        try:
            if path.lower().endswith('.xlsx') and Workbook:
                wb = Workbook()
                ws = wb.active
                ws.title = 'Exceptions'
                ws.append(headers)
                for row in exceptions:
                    ws.append(list(row))
                summary = wb.create_sheet('Summary')
                for row in [('Generated UTC', datetime.utcnow().isoformat()), ('GL rows', gl_count), ('Subledger rows', sub_count), ('Matched unique keys', matched), ('Exceptions', len(exceptions)), ('Note', 'Duplicate keys require review; no automatic netting or posting')]:
                    summary.append(row)
                wb.save(path)
            else:
                with open(path, 'w', newline='', encoding='utf-8-sig') as f:
                    writer = csv.writer(f)
                    writer.writerow(headers)
                    writer.writerows(exceptions)
            messagebox.showinfo('Export complete', f'Saved: {path}')
        except Exception as exc:
            messagebox.showerror('Export error', str(exc))

if __name__ == '__main__':
    App().mainloop()
